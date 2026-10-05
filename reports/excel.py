"""Libros Excel con los resultados numéricos."""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from calculos.resistencia import (LEYENDA_MARCAS, LEYENDA_POTENCIA, NOTAS, NUDO, avisos_casco,
                                  froudes, parametros, texto_rango_velocidad)


def _kilo(valor):
    """N -> kN o W -> kW; None si no está calculado."""
    return None if valor is None else valor / 1000


def _libro():
    libro = Workbook()
    libro.remove(libro.active)
    return libro


def _hoja(libro, titulo, cabecera, filas, formato="0.000", fijar="A2"):
    """Hoja con la cabecera en negrita y fija, y los números con 'formato'."""
    hoja = libro.create_sheet(titulo)
    hoja.append(cabecera)
    for fila in filas:
        hoja.append(list(fila))
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for columna in hoja.iter_cols(min_row=2):
        for celda in columna:
            if isinstance(celda.value, float):
                celda.number_format = formato
    for i, columna in enumerate(hoja.iter_cols(), start=1):
        ancho = max(len(f"{c.value:.3f}" if isinstance(c.value, float) else str(c.value or ""))
                    for c in columna)
        hoja.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 60)
    hoja.freeze_panes = fijar
    return hoja


def guardar_hidrostaticas(ruta, r):
    """hidrostaticas.xlsx: hidrostáticas en la DWL y datos de referencia."""
    libro = _libro()
    _hoja(libro, "Hidrostáticas", ["Magnitud", "Valor"], r["hidrostaticas"].items(), "0.0000")
    longitud, peso = r["unidades"]
    datos = [("Unidades de longitud", longitud), ("Unidades de peso", peso),
             ("Densidad [t/m3]", r["densidad"])]
    _hoja(libro, "Referencias", ["Dato", "Valor"], datos + list(r["referencias"].items()), "0.0000")
    libro.save(ruta)


def guardar_resistencia(ruta, r, v_max):
    """resistencia.xlsx: resumen a la velocidad de servicio, curvas por rendimiento,
    casco y validez."""
    casco, curvas = r["casco"], r["curvas"]
    rendimientos = list(curvas)
    metodos = list(curvas[rendimientos[0]])
    libro = _libro()

    cabecera = ["Método"]
    for eta in rendimientos:
        cabecera += [f"Rt [kN] η={eta:g}", f"P [kW] η={eta:g}"]
    cabecera += ["Marcas", "Corrección C_A", "Casco fuera de rango", "Velocidad fuera de rango", "Nota"]
    servicio = []
    for f in r["servicio"]:
        fila = [f["Método"]]
        for eta in rendimientos:
            fila += [_kilo(f["Rt"][eta]), _kilo(f["P"][eta])]
        servicio.append(fila + [f["Marcas"], f["Corrección"], f["Casco fuera de rango"],
                                f["Velocidad fuera de rango"], f["Nota"]])
    servicio += [("Marcas: " + LEYENDA_MARCAS,), ("Potencia: " + LEYENDA_POTENCIA,)]
    _hoja(libro, f"Servicio {r['v_servicio'] / NUDO:g} kn", cabecera, servicio)

    # Curvas: una hoja por magnitud y rendimiento; una fila por velocidad y una
    # columna por método (vacío = no calculado).
    velocidades = [p["v"] for p in curvas[rendimientos[0]][metodos[0]] if p["v"] <= v_max * (1 + 1e-9)]
    for eta in rendimientos:
        for clave, unidad in (("Rt", "kN"), ("P", "kW"), ("Rf", "kN"), ("Rr", "kN")):
            filas = []
            for i, v in enumerate(velocidades):
                fn = froudes(v, casco)
                filas.append([round(v / NUDO, 6), v, fn["FnL"], fn["Fnv"], fn["Fnb"]]
                             + [_kilo(curvas[eta][m][i][clave]) for m in metodos])
            _hoja(libro, f"{clave} {unidad} η={eta:g}",
                  ["V [kn]", "V [m/s]", "FnL", "Fnv", "Fnb"] + metodos, filas, fijar="B2")

    datos = [(k, ("codillo" if v else "pantoque redondo") if k == "Codillo" else v)
             for k, v in casco.items()]
    _hoja(libro, "Casco", ["Parámetro", "Valor"], datos + list(parametros(casco).items()), "0.0000")

    validez = [(m, NOTAS.get(m, ""), texto_rango_velocidad(m, casco),
                "; ".join(avisos_casco(m, casco)) or "dentro de rango") for m in metodos]
    _hoja(libro, "Validez", ["Método", "Nota", "Rango de velocidad", "Casco"], validez)
    libro.save(ruta)


def guardar_escalado(ruta, r):
    """escalado.xlsx: factores y control del escalado frente a la teoría."""
    libro = _libro()
    hoja = _hoja(libro, "Control", ["Magnitud", "Base", "Escalado", "Teórico", "Diferencia"],
                 r["control"]["filas"], "0.0000")
    for celda in hoja["E"][1:]:
        celda.number_format = "0.00%"
    f_eslora, f_manga, f_puntal, f_calado = r["factores"]
    calado_base, calado_escalado, calado_final = r["control"]["calados"]
    _hoja(libro, "Factores", ["Dato", "Valor"],
          [("Factor de eslora", f_eslora), ("Factor de manga", f_manga),
           ("Factor de puntal", f_puntal), ("Factor de calado", f_calado),
           ("Control (eslora, manga y puntal)", "correcto" if r["control"]["correcto"] else "NO CUADRA"),
           ("Calado base [m]", calado_base), ("Calado escalado con el puntal [m]", calado_escalado),
           ("Calado final [m]", calado_final)], "0.0000")
    libro.save(ruta)


def guardar_resumen_serie(ruta, resultados):
    """resumen_serie.xlsx: una fila por modelo escalado con hidrostáticas y
    resistencia y potencia a la velocidad de servicio para cada rendimiento."""
    primero = resultados[0]
    magnitudes = list(primero.get("hidrostaticas", {}))
    metodos = [f["Método"] for f in primero.get("servicio", [])]
    rendimientos = list(primero.get("curvas", {}))
    cabecera = ["Modelo", "F. eslora", "F. manga", "F. puntal", "F. calado"] + magnitudes
    for eta in rendimientos:
        for m in metodos:
            cabecera += [f"Rt {m} [kN] η={eta:g}", f"P {m} [kW] η={eta:g}"]

    filas = []
    for r in resultados:
        fila = [r["nombre"], *r["factores"]]
        fila += [r["hidrostaticas"][k] for k in magnitudes]
        for eta in rendimientos:
            for f in r["servicio"]:
                fila += [_kilo(f["Rt"][eta]), _kilo(f["P"][eta])]
        filas.append(fila)

    libro = _libro()
    _hoja(libro, "Serie", cabecera, filas, fijar="B2")
    libro.save(ruta)
